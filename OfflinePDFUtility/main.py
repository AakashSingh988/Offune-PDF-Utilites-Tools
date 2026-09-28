import gi
gi.require_version("Gtk", "3.0")
from gi.repository import Gtk

import tkinter as tk
from tkinter import ttk, messagebox
from pathlib import Path
import os
import subprocess
import pymupdf
from PIL import Image
from pypdf import PdfReader, PdfWriter
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from pptx import Presentation
from pptx.util import Inches


APP_DIR = Path("/OfflinePDFUtility")
OUTPUT_DIR = APP_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)


# =========================================================
# GENERAL HELPERS
# =========================================================

def choose_files(title, filetypes, multiple=False):
    dialog = Gtk.FileChooserDialog(
        title=title,
        action=Gtk.FileChooserAction.OPEN,
        buttons=(
            Gtk.STOCK_CANCEL,
            Gtk.ResponseType.CANCEL,
            Gtk.STOCK_OPEN,
            Gtk.ResponseType.OK
        )
    )

    dialog.set_select_multiple(multiple)

    for description, patterns in filetypes:
        file_filter = Gtk.FileFilter()
        file_filter.set_name(description)

        if isinstance(patterns, str):
            patterns = [patterns]

        for pattern in patterns:
            file_filter.add_pattern(pattern)

        dialog.add_filter(file_filter)

    files = []

    try:
        response = dialog.run()

        if response == Gtk.ResponseType.OK:
            files = dialog.get_filenames()

    finally:
        # Always close and destroy the GTK picker
        dialog.hide()
        dialog.destroy()

        # Process pending GTK events
        while Gtk.events_pending():
            Gtk.main_iteration()

    return files
def choose_output_file(title, extension, filetypes):
    dialog = Gtk.FileChooserDialog(
        title=title,
        action=Gtk.FileChooserAction.SAVE,
        buttons=(
            Gtk.STOCK_CANCEL,
            Gtk.ResponseType.CANCEL,
            Gtk.STOCK_SAVE,
            Gtk.ResponseType.OK
        )
    )

    dialog.set_do_overwrite_confirmation(True)

    for description, patterns in filetypes:
        file_filter = Gtk.FileFilter()
        file_filter.set_name(description)

        if isinstance(patterns, str):
            patterns = [patterns]

        for pattern in patterns:
            file_filter.add_pattern(pattern)

        dialog.add_filter(file_filter)

    output = None

    try:
        response = dialog.run()

        if response == Gtk.ResponseType.OK:
            output = dialog.get_filename()

            if output and not output.lower().endswith(
                extension.lower()
            ):
                output += extension

    finally:
        dialog.hide()
        dialog.destroy()

        while Gtk.events_pending():
            Gtk.main_iteration()

    return output
# =========================================================
# IMAGE → PDF
# =========================================================

def image_to_pdf():
    files = choose_files(
        "Select Images",
        [
            (
                "Image files",
                ["*.jpg", "*.jpeg", "*.png", "*.webp", "*.bmp"]
            )
        ],
        True
    )

    if not files:
        return

    output = choose_output_file(
        "Save PDF",
        ".pdf",
        [("PDF files", "*.pdf")]
    )

    if not output:
        return
    try:
        images = []

        for file in files:
            image = Image.open(file)

            if image.mode != "RGB":
                image = image.convert("RGB")

            images.append(image)

        images[0].save(
            output,
            save_all=True,
            append_images=images[1:]
        )

        for image in images:
            image.close()

        messagebox.showinfo(
            "Success",
            f"{len(files)} images converted into one PDF."
        )

    except Exception as e:
        messagebox.showerror(
            "Error",
            str(e)
        )

# =========================================================
# PDF → IMAGES
# =========================================================

def pdf_to_images():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    folder = choose_output_folder()

    if not folder:
        return

    try:
        document = pymupdf.open(pdf)

        for page_number, page in enumerate(document, start=1):
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
            output = os.path.join(
                folder,
                f"page_{page_number}.png"
            )
            pix.save(output)

        document.close()

        messagebox.showinfo(
            "Success",
            f"{len(document)} pages converted."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# OFFICE → PDF
# =========================================================

def office_to_pdf():
    files = choose_files(
        "Select Office files",
        [
            ("Word", "*.doc *.docx"),
            ("Excel", "*.xls *.xlsx"),
            ("PowerPoint", "*.ppt *.pptx"),
            ("Office files", "*.doc *.docx *.xls *.xlsx *.ppt *.pptx")
        ],
        True
    )

    if not files:
        return

    folder = choose_output_folder()

    if not folder:
        return

    try:
        for file in files:
            subprocess.run(
                [
                    "libreoffice",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    folder,
                    file
                ],
                check=True
            )

        messagebox.showinfo(
            "Success",
            "Office files converted to PDF."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# PDF → WORD
# =========================================================

def pdf_to_word():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    output = choose_output_file(
        "Save Word document",
        ".docx",
        [("Word document", "*.docx")]
    )

    if not output:
        return

    try:
        document = pymupdf.open(pdf)
        word = Document()

        for page in document:
            text = page.get_text()

            if text.strip():
                word.add_paragraph(text)
            else:
                word.add_paragraph("[No selectable text on this page]")

        word.save(output)
        document.close()

        messagebox.showinfo(
            "Success",
            "PDF converted to Word."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# PDF → EXCEL
# =========================================================

def pdf_to_excel():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    output = choose_output_file(
        "Save Excel file",
        ".xlsx",
        [("Excel workbook", "*.xlsx")]
    )

    if not output:
        return

    try:
        document = pymupdf.open(pdf)
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "PDF Text"

        row = 1

        for page_number, page in enumerate(document, start=1):
            sheet.cell(row, 1, f"Page {page_number}")
            row += 1

            text = page.get_text()

            for line in text.splitlines():
                sheet.cell(row, 1, line)
                row += 1

            row += 1

        workbook.save(output)
        document.close()

        messagebox.showinfo(
            "Success",
            "PDF text exported to Excel."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# PDF → POWERPOINT
# =========================================================

def pdf_to_ppt():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    output = choose_output_file(
        "Save PowerPoint",
        ".pptx",
        [("PowerPoint", "*.pptx")]
    )

    if not output:
        return

    temp_folder = APP_DIR / "output" / "_ppt_pages"
    temp_folder.mkdir(exist_ok=True)

    try:
        document = pymupdf.open(pdf)
        presentation = Presentation()

        blank_layout = presentation.slide_layouts[6]

        for number, page in enumerate(document, start=1):
            image_path = temp_folder / f"page_{number}.png"

            pix = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5))
            pix.save(str(image_path))

            slide = presentation.slides.add_slide(blank_layout)

            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height
            )

        presentation.save(output)
        document.close()

        messagebox.showinfo(
            "Success",
            "PDF converted to PowerPoint."

        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# MERGE PDFs
# =========================================================

def merge_pdfs():
    files = choose_files(
        "Select PDFs to merge",
        [("PDF files", "*.pdf")],
        True
    )

    if not files:
        return

    if len(files) < 2:
        messagebox.showwarning(
            "Merge PDF",
            "Please select at least 2 PDF files."
        )
        return

    output = choose_output_file(
        "Save merged PDF",
        ".pdf",
        [("PDF files", "*.pdf")]
    )

    if not output:
        return

    try:
        writer = PdfWriter()

        for file in files:
            reader = PdfReader(file)

            for page in reader.pages:
                writer.add_page(page)

        with open(output, "wb") as file:
            writer.write(file)

        messagebox.showinfo(
            "Success",
            f"{len(files)} PDFs merged successfully."
        )

    except Exception as e:
        messagebox.showerror(
            "Error",
            str(e)
        )

# =========================================================
# SPLIT PDF
# =========================================================

def split_pdf():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    folder = choose_output_folder()

    if not folder:
        return

    try:
        reader = PdfReader(pdf)

        for number, page in enumerate(reader.pages, start=1):
            writer = PdfWriter()
            writer.add_page(page)

            output = os.path.join(
                folder,
                f"page_{number}.pdf"
            )

            with open(output, "wb") as file:
                writer.write(file)

        messagebox.showinfo(
            "Success",
            f"{len(reader.pages)} pages split."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# ROTATE PDF
# =========================================================

def rotate_pdf():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    angle_window = tk.Toplevel(root)
    angle_window.title("Rotate PDF")
    angle_window.geometry("300x180")
    angle_window.resizable(False, False)

    tk.Label(
        angle_window,
        text="Choose rotation angle:"
    ).pack(pady=15)

    angle = tk.StringVar(value="90")

    combo = ttk.Combobox(
        angle_window,
        textvariable=angle,
        values=["90", "180", "270"],
        state="readonly"
    )

    combo.pack()

    def do_rotate():
        output = choose_output_file(
            "Save rotated PDF",
            ".pdf",
            [("PDF files", "*.pdf")]
        )

        if not output:
            return

        try:
            reader = PdfReader(pdf)
            writer = PdfWriter()

            for page in reader.pages:
                page.rotate(int(angle.get()))
                writer.add_page(page)

            with open(output, "wb") as file:
                writer.write(file)

            angle_window.destroy()

            messagebox.showinfo(
                "Success",
                "PDF rotated successfully."
            )

        except Exception as e:
            messagebox.showerror("Error", str(e))

    tk.Button(
        angle_window,
        text="Rotate",
        command=do_rotate,
        width=15
    ).pack(pady=20)


# =========================================================
# DELETE PAGES
# =========================================================

def delete_pages():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    page_window = tk.Toplevel(root)
    page_window.title("Delete Pages")
    page_window.geometry("400x220")

    tk.Label(
        page_window,
        text="Enter pages to delete\nExample: 1,3,5"
    ).pack(pady=15)

    entry = tk.Entry(page_window, width=35)
    entry.pack()

    def delete():
        try:
            pages = entry.get().replace(" ", "")
            delete_numbers = {
                int(x) for x in pages.split(",") if x
            }

            output = choose_output_file(
                "Save PDF",
                ".pdf",
                [("PDF files", "*.pdf")]
            )

            if not output:
                return

            reader = PdfReader(pdf)
            writer = PdfWriter()

            for number, page in enumerate(reader.pages, start=1):
                if number not in delete_numbers:
                    writer.add_page(page)

            with open(output, "wb") as file:
                writer.write(file)

            page_window.destroy()

            messagebox.showinfo(
                "Success",
                "Selected pages deleted."
            )

        except Exception as e:
            messagebox.showerror("Error", str(e))

    tk.Button(
        page_window,
        text="Delete Pages",
        command=delete,
        width=20
    ).pack(pady=20)


# =========================================================
# EXTRACT IMAGES
# =========================================================

def extract_images():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    folder = choose_output_folder()

    if not folder:
        return

    try:
        document = pymupdf.open(pdf)
        count = 0

        for page_number, page in enumerate(document, start=1):

            for image_number, image in enumerate(
                page.get_images(full=True),
                start=1
            ):

                xref = image[0]
                data = document.extract_image(xref)

                extension = data["ext"]

                output = os.path.join(
                    folder,
                    f"page_{page_number}_image_{image_number}.{extension}"
                )

                with open(output, "wb") as file:
                    file.write(data["image"])

                count += 1

        document.close()

        messagebox.showinfo(
            "Success",
            f"{count} images extracted."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# TEXT → PDF
# =========================================================

def text_to_pdf():
    text_file = choose_files(
        "Select text file",
        [("Text files", "*.txt"), ("All files", "*.*")],
        False
    )

    if not text_file:
        return

    output = choose_output_file(
        "Save PDF",
        ".pdf",
        [("PDF files", "*.pdf")]
    )

    if not output:
        return

    try:
        with open(
            text_file,
            "r",
            encoding="utf-8",
            errors="replace"
        ) as file:
            text = file.read()

        pdf = canvas.Canvas(output, pagesize=A4)

        width, height = A4
        y = height - 50

        for line in text.splitlines():

            if y < 50:
                pdf.showPage()
                y = height - 50

            pdf.drawString(40, y, line[:120])
            y -= 15

        pdf.save()

        messagebox.showinfo(
            "Success",
            "Text converted to PDF."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# COMPRESS PDF
# =========================================================

def compress_pdf():
    pdf = choose_files(
        "Select PDF",
        [("PDF files", "*.pdf")],
        False
    )

    if not pdf:
        return

    output = choose_output_file(
        "Save compressed PDF",
        ".pdf",
        [("PDF files", "*.pdf")]
    )

    if not output:
        return

    try:
        subprocess.run(
            [
                "gs",
                "-sDEVICE=pdfwrite",
                "-dCompatibilityLevel=1.4",
                "-dPDFSETTINGS=/ebook",
                "-dNOPAUSE",
                "-dQUIET",
                "-dBATCH",
                f"-sOutputFile={output}",
                pdf
            ],
            check=True
        )

        messagebox.showinfo(
            "Success",
            "PDF compressed successfully."
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))


# =========================================================
# OFFLINE PDF UTILITY - DASHBOARD UI
# =========================================================

root = tk.Tk()
root.title("Offline PDF Utility")
root.geometry("1150x720")
root.minsize(900, 600)

# =========================================================
# COLORS
# =========================================================

BG = "#0d1117"
SIDEBAR = "#161b22"
CARD = "#21262d"
CARD_HOVER = "#30363d"
TEXT = "#f0f6fc"
MUTED = "#8b949e"
ACCENT = "#58a6ff"
SUCCESS = "#3fb950"
BORDER = "#30363d"

root.configure(bg=BG)

# =========================================================
# TKINTER STYLE
# =========================================================

style = ttk.Style()
style.theme_use("clam")

style.configure(
    "App.TButton",
    font=("Arial", 10, "bold"),
    foreground=TEXT,
    background=CARD,
    borderwidth=0,
    padding=(14, 11)
)

style.map(
    "App.TButton",
    background=[
        ("active", CARD_HOVER),
        ("pressed", ACCENT)
    ]
)

style.configure(
    "Accent.TButton",
    font=("Arial", 10, "bold"),
    foreground="white",
    background=ACCENT,
    borderwidth=0,
    padding=(14, 11)
)

style.map(
    "Accent.TButton",
    background=[
        ("active", "#79b8ff"),
        ("pressed", ACCENT)
    ]
)

# =========================================================
# SIDEBAR
# =========================================================

sidebar = tk.Frame(
    root,
    bg=SIDEBAR,
    width=230
)

sidebar.pack(
    side="left",
    fill="y"
)

sidebar.pack_propagate(False)

# SIDEBAR LOGO
logo_frame = tk.Frame(
    sidebar,
    bg=SIDEBAR
)

logo_frame.pack(
    fill="x",
    padx=20,
    pady=(25, 30)
)

# Logo symbol
logo = tk.Label(
    logo_frame,
    text="◈",
    font=("Arial", 32, "bold"),
    fg=ACCENT,
    bg=SIDEBAR
)

logo.pack(
    side="left"
)

# Application name
logo_text = tk.Label(
    logo_frame,
    text="PDF\nUTILITY",
    font=("Arial", 16, "bold"),
    fg=TEXT,
    bg=SIDEBAR,
    justify="left"
)

logo_text.pack(
    side="left",
    padx=(10, 0)
)

# Creator name
creator = tk.Label(
    sidebar,
    text="Created by\nAakashSingh998",
    font=("Arial", 9),
    fg=MUTED,
    bg=SIDEBAR,
    justify="left"
)

creator.pack(
    anchor="w",
    padx=25,
    pady=(0, 20)
)

# Sidebar button creator

sidebar_buttons = {}

def sidebar_button(text, command):
    button = tk.Button(
        sidebar,
        text=text,
        command=command,
        anchor="w",
        font=("Arial", 11, "bold"),
        fg=MUTED,
        bg=SIDEBAR,
        activeforeground=TEXT,
        activebackground=CARD_HOVER,
        relief="flat",
        borderwidth=0,
        padx=25,
        pady=13,
        cursor="hand2"
    )

    button.pack(
        fill="x",
        pady=2
    )

    sidebar_buttons[text] = button

    return button

# =========================================================
# MAIN CONTENT
# =========================================================

content = tk.Frame(
    root,
    bg=BG
)

content.pack(
    side="left",
    fill="both",
    expand=True
)

# =========================================================
# SCROLLABLE CONTENT
# =========================================================

canvas = tk.Canvas(
    content,
    bg=BG,
    highlightthickness=0
)

scrollbar = ttk.Scrollbar(
    content,
    orient="vertical",
    command=canvas.yview
)

page = tk.Frame(
    canvas,
    bg=BG
)

page.bind(
    "<Configure>",
    lambda e: canvas.configure(
        scrollregion=canvas.bbox("all")
    )
)

canvas_window = canvas.create_window(
    (0, 0),
    window=page,
    anchor="nw"
)

canvas.configure(
    yscrollcommand=scrollbar.set
)

canvas.pack(
    side="left",
    fill="both",
    expand=True
)

scrollbar.pack(
    side="right",
    fill="y"
)

def resize_page(event):
    canvas.itemconfig(
        canvas_window,
        width=event.width
    )

canvas.bind(
    "<Configure>",
    resize_page
)

# =========================================================
# HEADER
# =========================================================

header = tk.Frame(
    page,
    bg=BG
)

header.pack(
    fill="x",
    padx=35,
    pady=(30, 10)
)

heading = tk.Label(
    header,
    text="Dashboard",
    font=("Arial", 26, "bold"),
    fg=TEXT,
    bg=BG
)

heading.pack(anchor="w")

description = tk.Label(
    header,
    text="Everything you need to work with PDF files — completely offline.",
    font=("Arial", 10),
    fg=MUTED,
    bg=BG
)

description.pack(
    anchor="w",
    pady=(5, 0)
)

# =========================================================
# STAT CARDS
# =========================================================

stats = tk.Frame(
    page,
    bg=BG
)

stats.pack(
    fill="x",
    padx=30,
    pady=20
)

def stat_card(parent, title, value, color):
    card = tk.Frame(
        parent,
        bg=CARD,
        highlightbackground=BORDER,
        highlightthickness=1
    )

    card.pack(
        side="left",
        fill="both",
        expand=True,
        padx=6
    )

    value_label = tk.Label(
        card,
        text=value,
        font=("Arial", 22, "bold"),
        fg=color,
        bg=CARD
    )

    value_label.pack(
        anchor="w",
        padx=18,
        pady=(15, 2)
    )

    title_label = tk.Label(
        card,
        text=title,
        font=("Arial", 9),
        fg=MUTED,
        bg=CARD
    )

    title_label.pack(
        anchor="w",
        padx=18,
        pady=(0, 15)
    )

stat_card(
    stats,
    "Available Tools",
    "13",
    ACCENT
)

stat_card(
    stats,
    "Internet Required",
    "NO",
    SUCCESS
)

stat_card(
    stats,
    "Multi-file Support",
    "YES",
    ACCENT
)

# =========================================================
# SECTION FUNCTION
# =========================================================

def section_title(parent, icon, title_text):
    frame = tk.Frame(
        parent,
        bg=BG
    )

    frame.pack(
        fill="x",
        padx=35,
        pady=(15, 8)
    )

    icon_label = tk.Label(
        frame,
        text=icon,
        font=("Arial", 17, "bold"),
        fg=ACCENT,
        bg=BG
    )

    icon_label.pack(
        side="left",
        padx=(0, 8)
    )

    title_label = tk.Label(
        frame,
        text=title_text,
        font=("Arial", 15, "bold"),
        fg=TEXT,
        bg=BG
    )

    title_label.pack(
        side="left"
    )

# =========================================================
# TOOL CARD
# =========================================================

def tool_card(parent, title, description, command, accent=False):
    card = tk.Frame(
        parent,
        bg=CARD,
        highlightbackground=BORDER,
        highlightthickness=1
    )

    card.pack(
        side="left",
        fill="both",
        expand=True,
        padx=6,
        pady=6
    )

    title_label = tk.Label(
        card,
        text=title,
        font=("Arial", 11, "bold"),
        fg=TEXT,
        bg=CARD
    )

    title_label.pack(
        anchor="w",
        padx=15,
        pady=(15, 4)
    )

    desc_label = tk.Label(
        card,
        text=description,
        font=("Arial", 9),
        fg=MUTED,
        bg=CARD,
        wraplength=210,
        justify="left"
    )

    desc_label.pack(
        anchor="w",
        padx=15,
        pady=(0, 10)
    )

    button = ttk.Button(
        card,
        text="Open Tool",
        command=command,
        style="Accent.TButton" if accent else "App.TButton"
    )

    button.pack(
        fill="x",
        padx=12,
        pady=(0, 12)
    )

    return card

# =========================================================
# CONVERSIONS
# =========================================================

section_title(
    page,
    "⇄",
    "Conversions"
)

conversion_container = tk.Frame(
    page,
    bg=BG
)

conversion_container.pack(
    fill="x",
    padx=29
)

row1 = tk.Frame(
    conversion_container,
    bg=BG
)

row1.pack(
    fill="x"
)

tool_card(
    row1,
    "Images → PDF",
    "Convert multiple JPG, PNG and other images into one PDF.",
    image_to_pdf,
    True
)

tool_card(
    row1,
    "PDF → Images",
    "Convert PDF pages into high-quality image files.",
    pdf_to_images
)

tool_card(
    row1,
    "Office → PDF",
    "Convert Word, Excel and PowerPoint files to PDF.",
    office_to_pdf
)

row2 = tk.Frame(
    conversion_container,
    bg=BG
)

row2.pack(
    fill="x"
)

tool_card(
    row2,
    "PDF → Word",
    "Extract PDF text into an editable Word document.",
    pdf_to_word
)

tool_card(
    row2,
    "PDF → Excel",
    "Extract PDF text into an Excel spreadsheet.",
    pdf_to_excel
)

tool_card(
    row2,
    "PDF → PowerPoint",
    "Convert PDF pages into PowerPoint slides.",
    pdf_to_ppt
)

# =========================================================
# PDF TOOLS
# =========================================================

section_title(
    page,
    "⚙",
    "PDF Tools"
)

tools_container = tk.Frame(
    page,
    bg=BG
)

tools_container.pack(
    fill="x",
    padx=29
)

tools_row1 = tk.Frame(
    tools_container,
    bg=BG
)

tools_row1.pack(
    fill="x"
)

tool_card(
    tools_row1,
    "Merge PDFs",
    "Combine two or more PDF files into one document.",
    merge_pdfs,
    True
)

tool_card(
    tools_row1,
    "Split PDF",
    "Separate a PDF into individual page files.",
    split_pdf
)

tool_card(
    tools_row1,
    "Rotate PDF",
    "Rotate PDF pages by 90, 180 or 270 degrees.",
    rotate_pdf
)

tools_row2 = tk.Frame(
    tools_container,
    bg=BG
)

tools_row2.pack(
    fill="x"
)

tool_card(
    tools_row2,
    "Delete Pages",
    "Remove selected pages from a PDF.",
    delete_pages
)

tool_card(
    tools_row2,
    "Compress PDF",
    "Reduce PDF file size using Ghostscript.",
    compress_pdf
)

tool_card(
    tools_row2,
    "Extract Images",
    "Extract embedded images from a PDF.",
    extract_images
)

tools_row3 = tk.Frame(
    tools_container,
    bg=BG
)

tools_row3.pack(
    fill="x"
)

tool_card(
    tools_row3,
    "Text → PDF",
    "Convert a text file into a PDF document.",
    text_to_pdf
)

# =========================================================
# STATUS BAR
# =========================================================

footer = tk.Frame(
    root,
    bg=SIDEBAR,
    height=42
)

footer.pack(
    side="bottom",
    fill="x"
)

status = tk.Label(
    footer,
    text="●  Ready",
    font=("Arial", 9),
    fg=SUCCESS,
    bg=SIDEBAR
)

status.pack(
    side="left",
    padx=20,
    pady=10
)

footer_text = tk.Label(
    footer,
    text="Offline PDF Utility  •  v1.0  •  AakashSingh998",    font=("Arial", 9),
    fg=MUTED,
    bg=SIDEBAR
)

footer_text.pack(
    side="right",
    padx=20
)

# =========================================================
# SIDEBAR ACTIONS
# =========================================================

def go_home():
    canvas.yview_moveto(0)

def go_conversions():
    canvas.yview_moveto(0.18)

def go_tools():
    canvas.yview_moveto(0.55)

def show_about():
    messagebox.showinfo(
        "About Offline PDF Utility",
        "Offline PDF Utility\n\n"
        "A local PDF conversion and management tool "
        "designed for Kali Linux.\n\n"
        "No cloud upload required.\n"
        "Your files remain on your computer."
    )

sidebar_button(
    "⌂   Home",
    go_home
)

sidebar_button(
    "⇄   Conversions",
    go_conversions
)

sidebar_button(
    "⚙   PDF Tools",
    go_tools
)

sidebar_button(
    "ⓘ   About",
    show_about
)

# =========================================================
# KEYBOARD SUPPORT
# =========================================================

root.bind(
    "<Escape>",
    lambda event: root.focus()
)

root.mainloop()
